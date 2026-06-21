#!/usr/bin/env python3
"""intent_gate.py — Pre-commit Intent Gate verification.

Reads docs/intent-trace.yaml and verifies that every business assertion
has a passing deterministic proof. Fires at HARDEN after verify.py.

Location: .agents/skills/verification-loop/scripts/intent_gate.py
Spec: docs/specs/intent-gate.md
"""

# ── PyYAML fail-closed ──────────────────────────────────────────
# PyYAML is a hard requirement. A stdlib regex fallback cannot guarantee
# hash equivalence with the YAML path — see spec "Design rationale".
import sys

try:
    import yaml
except ImportError:
    print("ERROR: PyYAML is required for Intent Gate hash verification.",
          file=sys.stderr)
    print("Install it: pip install pyyaml", file=sys.stderr)
    print("The Intent Gate sentinel exists (.sc-intent-gate-passed) but cannot",
          file=sys.stderr)
    print("be verified without PyYAML. Commit blocked (fail-closed).",
          file=sys.stderr)
    sys.exit(1)

import argparse
import hashlib
import json
import os
import re
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Optional


# ── Constants ───────────────────────────────────────────────────

DATA_RUNNER_ALLOWLIST = frozenset({"dbt", "great_expectations"})

# Patterns that indicate code/technical language in an assertion.
# Business assertions should be human-readable, not API contracts.
CODE_SYMBOL_PATTERNS = [
    re.compile(r"\b(?:GET|POST|PUT|PATCH|DELETE|HEAD|OPTIONS)\s+/"),
    re.compile(r"\b\d{3}\b"),  # HTTP status codes
    re.compile(r"\b(?:SELECT|INSERT|UPDATE|DELETE)\b", re.IGNORECASE),
    re.compile(r"\b(?:FROM|WHERE)\b"),  # SQL keywords case-sensitively to avoid matching prepositions
    re.compile(r"\bclass\s+\w+"),
    re.compile(r"\bdef\s+\w+"),
    re.compile(r"\bfunction\s+\w+"),
    re.compile(r"(?:->|=>)\s*\{"),
    re.compile(r"\w+\.\w+\("),  # method calls like foo.bar()
]

DEFAULT_TRACE_PATH = "docs/intent-trace.yaml"


# ── Parsing ─────────────────────────────────────────────────────

def parse_trace(trace_path: str) -> dict:
    """Parse and return trace YAML. Exit 1 on parse failure."""
    try:
        with open(trace_path) as f:
            data = yaml.safe_load(f)
    except yaml.YAMLError as exc:
        print(f"❌ Intent trace parse error: {exc}", file=sys.stderr)
        sys.exit(1)
    except OSError as exc:
        print(f"❌ Cannot read trace file: {exc}", file=sys.stderr)
        sys.exit(1)

    if not isinstance(data, dict):
        print("❌ Intent trace must be a YAML mapping", file=sys.stderr)
        sys.exit(1)

    return data


# ── Schema Validation ───────────────────────────────────────────

def validate_schema(trace: dict) -> list[str]:
    """Validate trace schema. Returns list of error messages (empty = valid)."""
    errors = []

    if "assertions" not in trace:
        errors.append("Missing required field: assertions")
        return errors

    assertions = trace["assertions"]
    if not isinstance(assertions, list):
        errors.append("'assertions' must be a list")
        return errors

    seen_ids = set()
    for i, assertion in enumerate(assertions):
        prefix = f"assertion[{i}]"

        if not isinstance(assertion, dict):
            errors.append(f"{prefix}: must be a mapping")
            continue

        # Required fields
        if "id" not in assertion:
            errors.append(f"{prefix}: missing required field 'id'")
            continue

        aid = assertion["id"]
        if aid in seen_ids:
            errors.append(f"{prefix}: duplicate id '{aid}'")
        seen_ids.add(aid)

        if "assertion" not in assertion:
            errors.append(f"{aid}: missing required field 'assertion'")

        if "proof" not in assertion:
            errors.append(f"{aid}: missing required field 'proof'")

        # Rationale check (AP-06)
        if "rationale" not in assertion:
            errors.append(f"{aid}: missing rationale field")
        elif assertion.get("assertion") and assertion["rationale"] == assertion["assertion"]:
            errors.append(
                f"{aid}: rationale repeats the assertion — "
                "must describe what the proof tests"
            )

    return errors


# ── Business Language Check ─────────────────────────────────────

def check_business_language(text: str) -> bool:
    """Check if assertion text is business language (not code/API).

    Returns True if the text passes the business language check.
    """
    for pattern in CODE_SYMBOL_PATTERNS:
        if pattern.search(text):
            return False
    return True


# ── Path Traversal Validation ───────────────────────────────────

def validate_trace_path(path: str, workspace: str) -> tuple[bool, str]:
    """Validate that a trace path is safe (relative, no traversal, inside workspace).

    Returns (is_valid, error_message).
    """
    if not path:
        return (False, "path is empty")
    if os.path.isabs(path):
        return (False, f"absolute path not allowed: {path}")
    if ".." in path.split(os.sep):
        return (False, f"path traversal not allowed: {path}")
    resolved = os.path.realpath(os.path.join(workspace, path))
    workspace_real = os.path.realpath(workspace)
    if not resolved.startswith(workspace_real + os.sep) and resolved != workspace_real:
        return (False, f"path escapes workspace: {path}")
    return (True, "")


def _validate_all_paths(trace: dict, workspace: str) -> list[str]:
    """Validate all paths in the trace. Returns list of errors."""
    errors = []
    for assertion in trace.get("assertions", []):
        aid = assertion.get("id", "?")
        proof = assertion.get("proof", {})

        # Validate reference path
        ref = proof.get("reference", "")
        if ref:
            valid, reason = validate_trace_path(ref, workspace)
            if not valid:
                errors.append(f"Unsafe path in trace ({aid}): {reason}")

        # Validate input path
        input_path = proof.get("input", "")
        if input_path:
            valid, reason = validate_trace_path(input_path, workspace)
            if not valid:
                errors.append(f"Unsafe path in trace ({aid}): {reason}")

        # Validate files
        for src in assertion.get("files", []):
            if src:
                valid, reason = validate_trace_path(src, workspace)
                if not valid:
                    errors.append(f"Unsafe path in trace ({aid}): {reason}")

    return errors


# ── UNTESTED Sign-off ───────────────────────────────────────────

def check_untested_signoff(assertion_id: str, workspace: str) -> bool:
    """Check if a human sign-off exists for an UNTESTED assertion."""
    signoff_path = os.path.join(workspace, f".sc-intent-untested-{assertion_id}")
    return os.path.exists(signoff_path)


# ── Proof Checks ────────────────────────────────────────────────

def check_code_proof(proof: dict, workspace: str) -> tuple[str, str]:
    """Run a test file and check exit code. Returns (status, reason)."""
    reference = proof.get("reference", "")
    if not reference:
        return ("FAIL", "proof missing 'reference' field")

    expected_exit = proof.get("expected_exit", 0)
    filepath = os.path.join(workspace, reference)

    if not os.path.exists(filepath):
        return ("FAIL", f"proof artifact not found: {reference}")

    ext = Path(reference).suffix
    runners = {
        ".py": [sys.executable, "-m", "pytest", reference, "-x", "--tb=short"],
        ".ts": ["npx", "--no-install", "vitest", "run", reference],
        ".tsx": ["npx", "--no-install", "vitest", "run", reference],
        ".js": ["npx", "--no-install", "vitest", "run", reference],
        ".go": ["go", "test", "-v", "-run", ".", f"./{Path(reference).parent}"],
        ".rs": ["cargo", "test", "--", "--test-threads=1"],
        ".java": ["./gradlew", "test", "--tests", Path(reference).stem],
        ".kt": ["./gradlew", "test", "--tests", Path(reference).stem],
    }
    cmd = runners.get(ext)
    if cmd is None:
        return ("FAIL", f"unsupported test file extension: {ext}")

    try:
        result = subprocess.run(cmd, cwd=workspace, capture_output=True, timeout=120)
    except FileNotFoundError:
        return ("FAIL", f"test runner not found for {ext}")
    except subprocess.TimeoutExpired:
        return ("FAIL", f"proof timed out after 120s: {reference}")

    if result.returncode == expected_exit:
        return ("PASS", "")
    return ("FAIL", f"test exited {result.returncode}, expected {expected_exit}")


def check_infra_proof(proof: dict, workspace: str) -> tuple[str, str]:
    """Evaluate an OPA/Rego policy against a plan/state file."""
    policy = os.path.join(workspace, proof.get("reference", ""))
    input_val = proof.get("input", "")
    if not input_val:
        return ("FAIL", "policy_eval requires a non-empty 'input' field")
    input_file = os.path.join(workspace, input_val)

    if not os.path.exists(policy):
        return ("FAIL", f"policy file not found: {proof.get('reference', '')}")

    # Try conftest first (exit code is reliable)
    try:
        cmd = ["conftest", "test", input_file, "--policy", policy, "--no-color"]
        result = subprocess.run(cmd, cwd=workspace, capture_output=True, timeout=60)
        if result.returncode == proof.get("expected_exit", 0):
            return ("PASS", "")
        stdout = result.stdout.decode(errors="replace").strip()
        stderr = result.stderr.decode(errors="replace").strip()
        return ("FAIL", f"conftest failed: {stdout or stderr}")
    except FileNotFoundError:
        pass

    # Fall back to opa eval — MUST parse JSON stdout, not trust exit code.
    try:
        cmd = ["opa", "eval", "--data", policy, "--input", input_file,
               "--format", "json", "data.main.deny"]
        result = subprocess.run(cmd, cwd=workspace, capture_output=True, timeout=60)
        if result.returncode != 0:
            stderr = result.stderr.decode(errors="replace").strip()
            return ("FAIL", f"opa eval command failed: {stderr}")

        try:
            output = json.loads(result.stdout)
            result_list = output.get("result", [])
            if not result_list:
                return ("PASS", "")

            expressions = result_list[0].get("expressions", [])
            violations = []
            for expr in expressions:
                val = expr.get("value")
                if val:
                    if isinstance(val, list):
                        violations.extend(str(v) for v in val)
                    else:
                        violations.append(str(val))

            if violations:
                return ("FAIL", f"policy violations: {', '.join(violations)}")
            return ("PASS", "")
        except (json.JSONDecodeError, KeyError, IndexError) as exc:
            return ("FAIL", f"failed to parse opa eval output: {exc}")
    except FileNotFoundError:
        pass

    return ("FAIL", "neither conftest nor opa found — install one to evaluate policies")


def check_data_proof(proof: dict, workspace: str) -> tuple[str, str]:
    """Run a data validation tool from the allowlist."""
    runner = proof.get("runner", "")
    if runner not in DATA_RUNNER_ALLOWLIST:
        return ("FAIL",
                f"data runner '{runner}' not in allowlist: "
                f"{sorted(DATA_RUNNER_ALLOWLIST)}")

    args = proof.get("args", [])
    if not isinstance(args, list):
        return ("FAIL", "proof.args must be a list, not a string")

    expected_exit = proof.get("expected_exit", 0)
    cmd = [runner] + [str(a) for a in args]

    try:
        result = subprocess.run(cmd, cwd=workspace, capture_output=True, timeout=300)
    except FileNotFoundError:
        return ("FAIL", f"data runner '{runner}' not found — install it")
    except subprocess.TimeoutExpired:
        return ("FAIL", f"proof timed out after 300s: {runner}")

    if result.returncode == expected_exit:
        return ("PASS", "")
    return ("FAIL", f"{runner} exited {result.returncode}, expected {expected_exit}")


def check_consulting_proof(proof: dict, workspace: str) -> tuple[str, str]:
    """Check sourcing traceability in a deliverable file."""
    reference = os.path.join(workspace, proof.get("reference", ""))
    if not os.path.exists(reference):
        return ("FAIL", f"deliverable not found: {proof.get('reference', '')}")

    pattern = re.compile(proof.get("source_pattern", r"\[SOURCE:[^\]]+\]"))
    min_sources = proof.get("min_sources", 1)

    ext = Path(reference).suffix.lower()
    if ext in (".pptx", ".pdf", ".docx"):
        try:
            text = _extract_text_from_binary(reference, ext)
        except ImportError as exc:
            return ("FAIL",
                    f"binary extraction library not installed ({ext}: {exc}). "
                    f"Cannot verify sourcing. Create .sc-intent-untested-<id> "
                    f"to accept without verification, or install the library.")
    else:
        text = Path(reference).read_text(encoding="utf-8", errors="ignore")

    matches = pattern.findall(text)
    if len(matches) >= min_sources:
        return ("PASS", "")
    return ("FAIL",
            f"found {len(matches)} source citations, "
            f"expected at least {min_sources}")


def _extract_text_from_binary(filepath: str, ext: str) -> str:
    """Extract text from binary document formats.

    Raises ImportError if the required library is not installed.
    """
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(filepath)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    parts.append(shape.text)
        return "\n".join(parts)

    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(filepath) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)

    if ext == ".docx":
        import docx
        doc = docx.Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs)

    return ""


# ── Proof dispatcher ────────────────────────────────────────────

PROOF_CHECKERS = {
    "test_exit_code": check_code_proof,
    "policy_eval": check_infra_proof,
    "data_runner": check_data_proof,
    "source_check": check_consulting_proof,
}


def check_proof(assertion: dict, workspace: str) -> tuple[str, str]:
    """Dispatch to the appropriate proof checker."""
    proof = assertion.get("proof", {})

    # Catch the common mistake: proof left as empty {} from STRATEGY
    if not proof or not proof.get("type"):
        return ("FAIL",
                "proof not linked — still empty from STRATEGY. "
                "Fill proof.type during SPEC "
                "(see spec-writing.md § 'Intent Trace — Proof Linking').")

    proof_type = proof["type"]
    checker = PROOF_CHECKERS.get(proof_type)
    if checker is None:
        return ("FAIL", f"unknown proof type: {proof_type}")

    return checker(proof, workspace)


# ── Dynamic Hash ────────────────────────────────────────────────

def compute_intent_hash(workspace: str, trace_path: str) -> Optional[str]:
    """Compute hash over trace + proof artifacts + source files.

    Hash covers:
    - The trace YAML itself
    - All proof artifact files (reference, input)
    - All source files listed in assertions[].files
    """
    trace_full = os.path.join(workspace, trace_path)
    if not os.path.exists(trace_full):
        return None

    paths_to_hash = [trace_path]

    with open(trace_full) as f:
        trace = yaml.safe_load(f)

    for assertion in trace.get("assertions", []):
        proof = assertion.get("proof", {})
        proof_type = proof.get("type", "")

        # Add proof reference (not for data_runner)
        if proof_type != "data_runner":
            ref = proof.get("reference", "")
            if ref and os.path.exists(os.path.join(workspace, ref)):
                paths_to_hash.append(ref)

        # Add input file for policy_eval
        input_file = proof.get("input", "")
        if input_file and os.path.exists(os.path.join(workspace, input_file)):
            paths_to_hash.append(input_file)

        # Add source files
        for src in assertion.get("files", []):
            if src and os.path.exists(os.path.join(workspace, src)):
                paths_to_hash.append(src)

    # Deduplicate, sort by byte order
    paths_to_hash = sorted(set(paths_to_hash), key=lambda p: p.encode())

    h = hashlib.sha256()
    for relpath in paths_to_hash:
        filepath = os.path.join(workspace, relpath)
        try:
            with open(filepath, "rb") as f:
                h.update(f.read())
        except OSError:
            continue

    return h.hexdigest()


# ── Sentinel ────────────────────────────────────────────────────

def write_intent_sentinel(workspace: str, trace_path: str, results: dict) -> str:
    """Write .sc-intent-gate-passed with dynamic hash."""
    sentinel_path = os.path.join(workspace, ".sc-intent-gate-passed")
    content_hash = compute_intent_hash(workspace, trace_path)

    with open(sentinel_path, "w") as f:
        f.write(f"done={datetime.now().isoformat()}\n")
        f.write(f"assertions={results['total']}\n")
        f.write(f"passed={results['passed']}\n")
        f.write(f"untested={results['untested']}\n")
        if content_hash:
            f.write(f"artifact_hash={content_hash}\n")

    return sentinel_path


# ── Main Gate Logic ─────────────────────────────────────────────

def run_gate(workspace: str, trace_rel: str = DEFAULT_TRACE_PATH,
             warn_if_missing: bool = False) -> dict:
    """Run the full Intent Gate. Returns result dict.

    Result dict:
        overall: "PASS" | "FAIL" | "SKIP"
        assertions: list of {id, assertion, status, reason}
        total: int
        passed: int
        failed: int
        untested: int
    """
    trace_path = os.path.join(workspace, trace_rel)

    # If trace doesn't exist → gate doesn't fire (backward compat)
    if not os.path.exists(trace_path):
        if warn_if_missing:
            blueprint = os.path.join(workspace, "docs", "strategic", "blueprint.md")
            if os.path.exists(blueprint):
                print(
                    "⚠️ WARNING: Full pipeline detected (blueprint.md exists) "
                    "but no intent trace.\n"
                    "Business assertions may have been skipped. Consider running:\n"
                    f"  intent_gate.py {workspace} --init\n"
                    "to create a trace from the blueprint's success metrics.",
                    file=sys.stderr,
                )
        return {"overall": "SKIP", "assertions": [],
                "total": 0, "passed": 0, "failed": 0, "untested": 0}

    # Parse
    trace = parse_trace(trace_path)

    # Validate schema
    schema_errors = validate_schema(trace)

    # Validate paths
    path_errors = _validate_all_paths(trace, workspace)

    all_errors = schema_errors + path_errors
    if all_errors:
        for err in all_errors:
            print(f"❌ {err}", file=sys.stderr)
        sys.exit(1)

    # Validate business language
    for assertion in trace.get("assertions", []):
        text = assertion.get("assertion", "")
        if not check_business_language(text):
            print(f"❌ Assertion {assertion['id']} contains code symbols: \"{text}\"",
                  file=sys.stderr)
            sys.exit(1)

    # Check proofs
    results = []
    for assertion in trace.get("assertions", []):
        aid = assertion["id"]
        assertion_text = assertion.get("assertion", "")
        scope = assertion.get("scope", "")

        # Ignore pre-filled status (AP-04)
        status, reason = check_proof(assertion, workspace)

        # If FAIL, check for UNTESTED sign-off
        if status == "FAIL" and check_untested_signoff(aid, workspace):
            status = "UNTESTED"
            reason = ""

        results.append({
            "id": aid,
            "assertion": assertion_text,
            "scope": scope,
            "status": status,
            "reason": reason,
        })

    passed = sum(1 for r in results if r["status"] == "PASS")
    failed = sum(1 for r in results if r["status"] == "FAIL")
    untested = sum(1 for r in results if r["status"] == "UNTESTED")
    overall = "PASS" if failed == 0 else "FAIL"

    return {
        "overall": overall,
        "assertions": results,
        "total": len(results),
        "passed": passed,
        "failed": failed,
        "untested": untested,
    }


# ── Verify hash only (commit gate) ─────────────────────────────

def verify_hash_only(workspace: str, trace_rel: str = DEFAULT_TRACE_PATH) -> bool:
    """Recompute hash and compare to sentinel. Returns True if match."""
    sentinel_path = os.path.join(workspace, ".sc-intent-gate-passed")
    if not os.path.exists(sentinel_path):
        return False

    # Read stored hash
    stored_hash = None
    with open(sentinel_path) as f:
        for line in f:
            if line.startswith("artifact_hash="):
                stored_hash = line.strip().split("=", 1)[1]
                break

    if stored_hash is None:
        # No hash in sentinel — presence-only mode, pass
        return True

    computed = compute_intent_hash(workspace, trace_rel)
    if computed is None:
        return False

    return computed == stored_hash


# ── CLI ─────────────────────────────────────────────────────────

def _format_human_report(result: dict) -> str:
    """Format gate result as human-readable report."""
    lines = []
    lines.append(f"Intent Gate: {result['overall']}")
    lines.append(f"  Total: {result['total']}  "
                 f"Passed: {result['passed']}  "
                 f"Failed: {result['failed']}  "
                 f"Untested: {result['untested']}")
    lines.append("")

    for a in result["assertions"]:
        icon = {"PASS": "✅", "FAIL": "❌", "UNTESTED": "⚠️"}.get(a["status"], "?")
        lines.append(f"  {icon} {a['id']}: {a['assertion']}")
        if a["reason"]:
            lines.append(f"     Reason: {a['reason']}")

    # Consulting scope caveat — must derive from actual assertion scopes
    has_consulting = any(a.get("scope") == "consulting"
                         for a in result.get("assertions", []))
    if has_consulting:
        lines.append("")
        lines.append("  ⚠️ Consulting scope: sourcing traceability only. "
                      "This verifies that claims cite sources, NOT that claims are correct.")

    return "\n".join(lines)


# ── Correction Signal Integration ───────────────────────────────

def _emit_intent_signal(result: dict, workspace: str) -> None:
    """Emit a correction signal for the first failing intent assertion."""
    try:
        # Script Resolution Order: project path → plugin fallback
        scripts_dir = str(Path(__file__).resolve().parent)
        if scripts_dir not in sys.path:
            sys.path.insert(0, scripts_dir)
        from correction_signal import (
            create_signal, write_signal, LocalizationData,
        )
    except ImportError:
        try:
            plugin_dir = str(Path.home() / ".gemini" / "config" / "plugins" / "stream-coding" / "skills" / "verification-loop" / "scripts")
            sys.path.insert(0, plugin_dir)
            from correction_signal import (
                create_signal, write_signal, LocalizationData,
            )
        except ImportError:
            return  # SCRIPT_NOT_FOUND — graceful degradation

    failing = [a for a in result.get("assertions", []) if a["status"] == "FAIL"]
    if not failing:
        return

    first_fail = failing[0]
    aid = first_fail["id"]

    # Extract files from the assertion scope
    files = first_fail.get("files", [])
    if not files:
        files = []

    localization = LocalizationData(
        files=files,
        lines=[],
        assertion_id=aid,
        finding=f"Intent assertion {aid} FAIL: {first_fail.get('reason', '')}",
        tool_output=json.dumps(result, indent=2)[:2000],
    )

    signal = create_signal(
        gate="intent",
        check=aid,
        localization=localization,
        workspace=workspace,
        sentinel_name="intent-gate-passed",
    )
    write_signal(signal, workspace)


def _clear_intent_signal(workspace: str) -> None:
    """Clear any existing intent signal on PASS."""
    try:
        scripts_dir = str(Path(__file__).resolve().parent)
        if scripts_dir not in sys.path:
            sys.path.insert(0, scripts_dir)
        from correction_signal import clear_signal
    except ImportError:
        try:
            plugin_dir = str(Path.home() / ".gemini" / "config" / "plugins" / "stream-coding" / "skills" / "verification-loop" / "scripts")
            sys.path.insert(0, plugin_dir)
            from correction_signal import clear_signal
        except ImportError:
            return
    clear_signal("intent", workspace)


def main():
    parser = argparse.ArgumentParser(description="Intent Gate — verify business assertions")
    parser.add_argument("workspace", help="Project root directory")
    parser.add_argument("--trace", default=DEFAULT_TRACE_PATH,
                        help=f"Path to trace YAML (default: {DEFAULT_TRACE_PATH})")
    parser.add_argument("--json", action="store_true", dest="json_output",
                        help="Output results as JSON")
    parser.add_argument("--verify-hash-only", action="store_true",
                        help="Only verify sentinel hash (for commit gate)")
    parser.add_argument("--warn-if-missing", action="store_true",
                        help="Warn if no trace exists in a full pipeline")

    args = parser.parse_args()

    if args.verify_hash_only:
        if verify_hash_only(args.workspace, args.trace):
            sys.exit(0)
        else:
            print("BLOCKED: Intent Gate hash verification failed",
                  file=sys.stderr)
            sys.exit(1)

    result = run_gate(args.workspace, args.trace, args.warn_if_missing)

    if result["overall"] == "SKIP":
        sys.exit(0)

    if args.json_output:
        print(json.dumps(result, indent=2))
    else:
        print(_format_human_report(result))

    if result["overall"] == "PASS":
        _clear_intent_signal(args.workspace)
        trace_path = args.trace
        sentinel = write_intent_sentinel(args.workspace, trace_path, result)
        if not args.json_output:
            print(f"\n  Sentinel written: {sentinel}")
        else:
            print(f"Sentinel written: {sentinel}", file=sys.stderr)
        sys.exit(0)
    else:
        _emit_intent_signal(result, args.workspace)
        sys.exit(1)


if __name__ == "__main__":
    main()
